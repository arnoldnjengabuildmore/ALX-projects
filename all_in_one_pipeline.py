"""
all_in_one_pipeline.py

Single entry-point pipeline that creates one workbook with these sheets in order:
1) Raw Data 1
2) TMC
3) flow profile
4) Summary

It orchestrates existing scripts:
- Final Calculation Flow.py
- flow_profile_pipeline.py
- peak_period_summary.py

Usage:
    python all_in_one_pipeline.py --input "Raw Data 1.xlsx" --output "Combined Report.xlsx"
"""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import time
from copy import copy
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox
except Exception:
    tk = None
    filedialog = None
    messagebox = None


def run_cmd(cmd: list[str], cwd: Path) -> None:
    result = subprocess.run(cmd, cwd=str(cwd), text=True, capture_output=True)
    if result.returncode != 0:
        msg = [
            f"Command failed: {' '.join(cmd)}",
            f"stdout:\n{result.stdout.strip()}",
            f"stderr:\n{result.stderr.strip()}",
        ]
        raise RuntimeError("\n".join(msg))


def default_report_name_for_input(input_path: Path) -> str:
    stem = input_path.stem.strip()
    lowered = stem.lower()
    if lowered.startswith("raw "):
        stem = stem[4:].strip() or input_path.stem.strip()
    return f"Final_Report_{stem}.xlsx"


def ensure_io_folders(script_root: Path) -> tuple[Path, Path]:
    input_dir = script_root / "input"
    output_dir = script_root / "output"
    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    return input_dir, output_dir


def prepare_default_input_file(script_root: Path, input_dir: Path) -> Path:
    default_input = input_dir / "Raw Data 1.xlsx"
    legacy_input = script_root / "Raw Data 1.xlsx"
    if not default_input.exists() and legacy_input.exists():
        shutil.copy2(legacy_input, default_input)
    return default_input


def newest_processed_file(folder: Path, after_ts: float) -> Optional[Path]:
    candidates = []
    for f in folder.glob("*_processed_*.xlsx"):
        if f.name.startswith("~$"):
            continue
        try:
            mt = f.stat().st_mtime
        except OSError:
            continue
        if mt >= after_ts:
            candidates.append((mt, f))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0][1]


def copy_sheet(src_ws, dst_ws) -> None:
    # Cell values and styles
    for row in src_ws.iter_rows(min_row=1, max_row=src_ws.max_row, min_col=1, max_col=src_ws.max_column):
        for src_cell in row:
            dst_cell = dst_ws.cell(row=src_cell.row, column=src_cell.column, value=src_cell.value)
            if src_cell.has_style:
                dst_cell.font = copy(src_cell.font)
                dst_cell.fill = copy(src_cell.fill)
                dst_cell.border = copy(src_cell.border)
                dst_cell.alignment = copy(src_cell.alignment)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy(src_cell.protection)
            if src_cell.hyperlink:
                dst_cell._hyperlink = copy(src_cell.hyperlink)
            if src_cell.comment:
                dst_cell.comment = copy(src_cell.comment)

    # Column widths
    for key, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width
        dst_ws.column_dimensions[key].hidden = dim.hidden

    # Row heights
    for idx, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[idx].height = dim.height
        dst_ws.row_dimensions[idx].hidden = dim.hidden

    # Merged ranges
    for merged in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged))

    # Basic sheet properties
    dst_ws.freeze_panes = src_ws.freeze_panes
    dst_ws.sheet_view.zoomScale = src_ws.sheet_view.zoomScale


def ensure_sheet_order_and_names(final_path: Path, raw_input: Path) -> None:
    wb = load_workbook(final_path)

    def _find_sheet(*aliases: str) -> Optional[str]:
        want = {a.strip().lower() for a in aliases}
        for sheet_name in wb.sheetnames:
            if sheet_name.strip().lower() in want:
                return sheet_name
        return None

    flow_name = _find_sheet("Flow Profile", "flow profile")
    summary_name = _find_sheet("Peak Period Summary", "Summary")
    raw_existing = _find_sheet("Raw Data 1")

    def _rename_exact(current_name: str, target_name: str) -> str:
        if current_name == target_name:
            return current_name
        current_lower = current_name.strip().lower()
        target_lower = target_name.strip().lower()
        if current_lower == target_lower:
            tmp = target_name + "__tmp__"
            wb[current_name].title = tmp
            wb[tmp].title = target_name
            return target_name
        wb[current_name].title = target_name
        return target_name

    tmc_candidate = None
    for n in wb.sheetnames:
        if n in (flow_name, summary_name, raw_existing):
            continue
        tmc_candidate = n
        break
    if tmc_candidate is None:
        raise ValueError("Could not locate TMC/source sheet in consolidated workbook.")

    # Ensure Raw Data 1 exists as first sheet.
    if raw_existing and raw_existing in wb.sheetnames:
        del wb[raw_existing]
    raw_wb = load_workbook(raw_input)
    raw_src = raw_wb[raw_wb.sheetnames[0]]
    raw_dst = wb.create_sheet("Raw Data 1", 0)
    copy_sheet(raw_src, raw_dst)

    # Rename required sheets in-place so charts/objects are preserved.
    tmc_name = _rename_exact(tmc_candidate, "TMC")
    if flow_name and flow_name in wb.sheetnames:
        flow_name = _rename_exact(flow_name, "flow profile")
    else:
        flow_name = wb.create_sheet("flow profile").title
    if summary_name and summary_name in wb.sheetnames:
        summary_name = _rename_exact(summary_name, "Summary")
    else:
        summary_name = wb.create_sheet("Summary").title

    # Keep only required sheets and enforce order.
    required = ["Raw Data 1", "TMC", "flow profile", "Summary"]
    required_lower = {name.lower() for name in required}
    for name in list(wb.sheetnames):
        if name.lower() not in required_lower:
            del wb[name]
    name_map = {name.lower(): name for name in wb.sheetnames}
    wb._sheets = [wb[name_map[name.lower()]] for name in required if name.lower() in name_map]

    raw_wb.close()
    wb.save(final_path)
    wb.close()


def build_report(input_path: Path, output_path: Path, *, show_success_dialog: bool = False) -> Path:
    cwd = Path.cwd()
    script_root = Path(__file__).resolve().parent
    pipeline_dir = script_root / "pipeline_core"
    _, output_dir = ensure_io_folders(script_root)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not input_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")

    final_calc = pipeline_dir / "Final Calculation Flow.py"
    flow_profile = pipeline_dir / "flow_profile_pipeline.py"
    summary = pipeline_dir / "peak_period_summary.py"

    for script in (final_calc, flow_profile, summary):
        if not script.exists():
            raise FileNotFoundError(f"Required script missing: {script}")

    # Final Calculation Flow.py is hardcoded to read `Raw Data 1.xlsx`
    fixed_input_name = cwd / "Raw Data 1.xlsx"
    backup_path = None
    if fixed_input_name.exists():
        same_file = fixed_input_name.resolve() == input_path
        if not same_file:
            backup_path = cwd / f"Raw Data 1.xlsx.backup_{int(time.time())}"
            shutil.copy2(fixed_input_name, backup_path)
            shutil.copy2(input_path, fixed_input_name)
    else:
        shutil.copy2(input_path, fixed_input_name)

    try:
        temp_files: list[Path] = []
        started = time.time()
        run_cmd([sys.executable, str(final_calc)], cwd)
        processed = newest_processed_file(cwd, started)
        if processed is None:
            # Fallback to newest processed file overall.
            processed = newest_processed_file(cwd, 0)
        if processed is None:
            raise RuntimeError("Could not detect processed workbook output from Final Calculation Flow.py")
        temp_files.append(processed)

        flow_out = output_dir / f"{output_path.stem}_tmp_flow_profile.xlsx"
        temp_files.append(flow_out)
        run_cmd(
            [sys.executable, str(flow_profile), "--processed", str(processed), "--output", str(flow_out)],
            cwd,
        )

        summary_out = output_dir / f"{output_path.stem}_tmp_summary.xlsx"
        temp_files.append(summary_out)
        run_cmd(
            [sys.executable, str(summary), "--processed", str(flow_out), "--output", str(summary_out)],
            cwd,
        )

        ensure_sheet_order_and_names(summary_out, input_path)

        try:
            if output_path.exists():
                output_path.unlink(missing_ok=True)
            shutil.move(str(summary_out), str(output_path))
        except PermissionError as exc:
            raise RuntimeError(
                f"Cannot write final report because the file is open: {output_path}. "
                "Please close it in Excel/Explorer preview and run again."
            ) from exc

        # Keep only the final combined report.
        for temp_path in temp_files:
            if temp_path.exists() and temp_path.resolve() != output_path.resolve():
                temp_path.unlink(missing_ok=True)

        print(f"Combined workbook written to: {output_path}")
        if show_success_dialog and messagebox is not None and tk is not None:
            root = tk.Tk()
            root.withdraw()
            messagebox.showinfo("Success", "File has been produced.\nThank you")
            root.destroy()
    finally:
        if backup_path and backup_path.exists():
            shutil.copy2(backup_path, fixed_input_name)
            backup_path.unlink(missing_ok=True)

    return output_path


def launch_gui(default_input: Path, default_output: Path) -> None:
    if tk is None or filedialog is None or messagebox is None:
        raise RuntimeError("Tkinter is not available in this Python environment.")

    root = tk.Tk()
    root.title("Traffic Combined Report")
    root.geometry("700x210")
    root.resizable(False, False)

    input_var = tk.StringVar(value=str(default_input))
    output_var = tk.StringVar(value=str(default_output))
    status_var = tk.StringVar(value="Ready")

    def choose_input() -> None:
        path = filedialog.askopenfilename(
            title="Select Raw Data Workbook",
            filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")],
            initialdir=str(default_input.parent),
        )
        if path:
            chosen = Path(path)
            input_var.set(path)
            output_var.set(str(chosen.with_name(default_report_name_for_input(chosen))))

    def choose_output() -> None:
        path = filedialog.asksaveasfilename(
            title="Save Final Report As",
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")],
            initialdir=str(default_output.parent),
            initialfile=default_output.name,
        )
        if path:
            output_var.set(path)

    def run_pipeline() -> None:
        run_button.config(state="disabled")
        status_var.set("Running pipeline... please wait")
        root.update_idletasks()
        try:
            out_path = build_report(
                Path(input_var.get()).expanduser().resolve(),
                Path(output_var.get()).expanduser().resolve(),
                show_success_dialog=False,
            )
            status_var.set(f"Done: {out_path.name}")
            messagebox.showinfo("Success", "File has been produced.\nThank you")
        except Exception as exc:
            status_var.set("Failed")
            messagebox.showerror("Error", str(exc))
        finally:
            run_button.config(state="normal")

    pad_x = 12
    pad_y = 10
    tk.Label(root, text="Input (Raw Data .xlsx):").grid(row=0, column=0, sticky="w", padx=pad_x, pady=(pad_y, 4))
    tk.Entry(root, textvariable=input_var, width=72).grid(row=1, column=0, padx=pad_x, pady=0, sticky="w")
    tk.Button(root, text="Browse", width=10, command=choose_input).grid(row=1, column=1, padx=(0, pad_x), pady=0)

    tk.Label(root, text="Output (Final report):").grid(row=2, column=0, sticky="w", padx=pad_x, pady=(pad_y, 4))
    tk.Entry(root, textvariable=output_var, width=72).grid(row=3, column=0, padx=pad_x, pady=0, sticky="w")
    tk.Button(root, text="Browse", width=10, command=choose_output).grid(row=3, column=1, padx=(0, pad_x), pady=0)

    run_button = tk.Button(root, text="Generate Final Report", width=24, command=run_pipeline)
    run_button.grid(row=4, column=0, sticky="w", padx=pad_x, pady=(16, 0))
    tk.Label(root, textvariable=status_var, fg="#1f4e79").grid(row=5, column=0, sticky="w", padx=pad_x, pady=(10, 0))

    root.mainloop()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", "-i", default="Raw Data 1.xlsx", help="Raw input workbook")
    parser.add_argument("--output", "-o", default=None, help="Final workbook output path")
    parser.add_argument("--gui", action="store_true", help="Launch Tkinter GUI")
    parser.add_argument("--no-gui", action="store_true", help="Suppress Tkinter dialogs")
    args = parser.parse_args()

    script_root = Path(__file__).resolve().parent
    input_dir, output_dir = ensure_io_folders(script_root)

    default_input = prepare_default_input_file(script_root, input_dir)

    if args.input:
        arg_input = Path(args.input)
        if arg_input.is_absolute():
            input_path = arg_input
        elif arg_input.exists():
            input_path = arg_input.resolve()
        else:
            input_path = (input_dir / arg_input.name).resolve()
    else:
        input_path = default_input.resolve()

    if args.output:
        arg_output = Path(args.output)
        if arg_output.is_absolute():
            output_path = arg_output
        elif arg_output.parent != Path('.'):
            output_path = (Path.cwd() / arg_output).resolve()
        else:
            output_path = (output_dir / arg_output.name).resolve()
    else:
        output_path = (output_dir / default_report_name_for_input(input_path)).resolve()

    launch_window = args.gui or (len(sys.argv) == 1 and tk is not None)
    if launch_window:
        launch_gui(input_path, output_path)
        return

    build_report(input_path, output_path, show_success_dialog=not args.no_gui)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        if messagebox is not None and tk is not None:
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror("Error", str(exc))
            root.destroy()
        raise
